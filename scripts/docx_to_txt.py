"""
Word (.docx) → txt
用法: python scripts/docx_to_txt.py <docx路径>
输出: 同目录的 <文件名>.txt
"""
import os
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document


def main():
    if len(sys.argv) < 2:
        print("用法: python scripts/docx_to_txt.py <docx路径>")
        sys.exit(1)

    docx_path = sys.argv[1]
    if not os.path.exists(docx_path):
        print(f"❌ 文件不存在: {docx_path}")
        sys.exit(1)

    print(f"📄 处理: {docx_path}")
    doc = Document(docx_path)

    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    # 也处理表格里的内容（如果有）
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in lines:
                    lines.append(text)

    full_text = "\n".join(lines)

    out_path = os.path.splitext(docx_path)[0] + ".txt"
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(full_text)

    print(f"✅ 完成！输出: {out_path}")
    print(f"📊 总字符数: {len(full_text)}")
    print(f"📝 段落数: {len(lines)}")


if __name__ == "__main__":
    main()